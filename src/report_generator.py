"""
Report Generator Module

Generates CSV and Excel reports from detection results.
"""

import csv
import logging
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

from .thin_detector import DetectionResult, ThinPageStatus, summarize_results

logger = logging.getLogger(__name__)


class ReportGenerator:
    """Generates reports from thin page detection results."""

    # Column configuration
    COLUMNS = [
        'URL',
        'Language',
        'SKU',
        'Page Title',
        'Word Count',
        'Has Description',
        'Has Specifications',
        'Image Count',
        'Thin Page Flag',
        'Reason',
        'Timestamp'
    ]

    def __init__(self, output_dir: str = './output'):
        """
        Initialize the report generator.

        Args:
            output_dir: Directory for output files
        """
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def _generate_filename(self, extension: str) -> str:
        """Generate timestamped filename."""
        timestamp = datetime.now().strftime('%Y-%m-%d_%H%M%S')
        return f"thin_pages_report_{timestamp}.{extension}"

    def _results_to_dataframe(
        self,
        results: List[DetectionResult],
        timestamp: Optional[datetime] = None
    ) -> pd.DataFrame:
        """Convert detection results to a pandas DataFrame."""
        if timestamp is None:
            timestamp = datetime.now()

        timestamp_str = timestamp.strftime('%Y-%m-%d %H:%M:%S')

        data = []
        for result in results:
            data.append({
                'URL': result.url,
                'Language': result.language,
                'SKU': result.sku,
                'Page Title': result.page_title,
                'Word Count': result.word_count,
                'Has Description': 'Yes' if result.has_description else 'No',
                'Has Specifications': 'Yes' if result.has_specifications else 'No',
                'Image Count': result.image_count,
                'Thin Page Flag': result.status.value,
                'Reason': '; '.join(result.reasons) if result.reasons else (result.error or ''),
                'Timestamp': timestamp_str
            })

        return pd.DataFrame(data, columns=self.COLUMNS)

    def generate_csv(
        self,
        results: List[DetectionResult],
        filename: Optional[str] = None
    ) -> str:
        """
        Generate CSV report.

        Args:
            results: List of detection results
            filename: Optional custom filename

        Returns:
            Path to generated CSV file
        """
        if filename is None:
            filename = self._generate_filename('csv')

        filepath = self.output_dir / filename

        df = self._results_to_dataframe(results)
        df.to_csv(filepath, index=False, quoting=csv.QUOTE_ALL)

        logger.info(f"Generated CSV report: {filepath}")
        return str(filepath)

    def generate_excel(
        self,
        results: List[DetectionResult],
        filename: Optional[str] = None,
        include_summary: bool = True
    ) -> str:
        """
        Generate Excel report with formatting.

        Args:
            results: List of detection results
            filename: Optional custom filename
            include_summary: Include summary statistics sheet

        Returns:
            Path to generated Excel file
        """
        if filename is None:
            filename = self._generate_filename('xlsx')

        filepath = self.output_dir / filename
        timestamp = datetime.now()

        df = self._results_to_dataframe(results, timestamp)

        # Create workbook
        wb = Workbook()

        # Main data sheet
        ws_data = wb.active
        ws_data.title = "Thin Pages Report"

        # Write data with formatting
        self._write_data_sheet(ws_data, df)

        # Summary sheet
        if include_summary:
            ws_summary = wb.create_sheet("Summary")
            summary = summarize_results(results)
            self._write_summary_sheet(ws_summary, summary, timestamp)

        # Thin pages only sheet
        thin_results = [r for r in results if r.status == ThinPageStatus.THIN]
        if thin_results:
            ws_thin = wb.create_sheet("Thin Pages Only")
            df_thin = self._results_to_dataframe(thin_results, timestamp)
            self._write_data_sheet(ws_thin, df_thin)

        # Error pages sheet
        error_results = [r for r in results if r.status == ThinPageStatus.ERROR]
        if error_results:
            ws_errors = wb.create_sheet("Errors")
            df_errors = self._results_to_dataframe(error_results, timestamp)
            self._write_data_sheet(ws_errors, df_errors)

        wb.save(filepath)
        logger.info(f"Generated Excel report: {filepath}")
        return str(filepath)

    def _write_data_sheet(self, ws, df: pd.DataFrame):
        """Write data to worksheet with formatting."""
        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        thin_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        ok_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        error_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Write headers
        for col_num, column in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=col_num, value=column)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border

        # Write data rows
        for row_num, row in enumerate(df.itertuples(index=False), 2):
            for col_num, value in enumerate(row, 1):
                cell = ws.cell(row=row_num, column=col_num, value=value)
                cell.border = border

                # Apply conditional formatting based on status
                if col_num == 9:  # Thin Page Flag column
                    if value == "THIN":
                        cell.fill = thin_fill
                    elif value == "OK":
                        cell.fill = ok_fill
                    elif value == "ERROR":
                        cell.fill = error_fill

        # Auto-adjust column widths
        for col_num, column in enumerate(df.columns, 1):
            max_length = len(str(column))
            for row in df.itertuples(index=False):
                cell_value = str(row[col_num - 1])
                max_length = max(max_length, min(len(cell_value), 50))

            adjusted_width = max_length + 2
            ws.column_dimensions[chr(64 + col_num) if col_num <= 26 else 'A' + chr(64 + col_num - 26)].width = adjusted_width

        # Freeze header row
        ws.freeze_panes = 'A2'

    def _write_summary_sheet(self, ws, summary: Dict[str, Any], timestamp: datetime):
        """Write summary statistics to worksheet."""
        # Styles
        title_font = Font(bold=True, size=14)
        header_font = Font(bold=True)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        row = 1

        # Title
        ws.cell(row=row, column=1, value="Thin Pages Analysis Summary").font = title_font
        row += 2

        # Timestamp
        ws.cell(row=row, column=1, value="Analysis Date:")
        ws.cell(row=row, column=2, value=timestamp.strftime('%Y-%m-%d %H:%M:%S'))
        row += 2

        # Overall statistics
        ws.cell(row=row, column=1, value="Overall Statistics").font = header_font
        row += 1

        stats = [
            ("Total URLs Analyzed", summary['total']),
            ("Thin Pages Found", f"{summary['thin']} ({summary['thin_percent']}%)"),
            ("Acceptable Pages", f"{summary['ok']} ({summary['ok_percent']}%)"),
            ("Failed URLs", f"{summary['error']} ({summary['error_percent']}%)"),
        ]

        for label, value in stats:
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=value)
            row += 1

        row += 1

        # Reason breakdown
        if summary.get('reason_breakdown'):
            ws.cell(row=row, column=1, value="Breakdown by Reason").font = header_font
            row += 1

            for reason, count in sorted(summary['reason_breakdown'].items(), key=lambda x: -x[1]):
                ws.cell(row=row, column=1, value=reason)
                ws.cell(row=row, column=2, value=count)
                row += 1

            row += 1

        # Language breakdown
        if summary.get('language_breakdown'):
            ws.cell(row=row, column=1, value="Breakdown by Language").font = header_font
            row += 1

            # Headers
            headers = ['Language', 'Total', 'Thin', 'OK', 'Error']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col, value=header)
                cell.font = header_font
                cell.border = border
            row += 1

            for lang, stats in sorted(summary['language_breakdown'].items()):
                ws.cell(row=row, column=1, value=lang).border = border
                ws.cell(row=row, column=2, value=stats['total']).border = border
                ws.cell(row=row, column=3, value=stats['thin']).border = border
                ws.cell(row=row, column=4, value=stats['ok']).border = border
                ws.cell(row=row, column=5, value=stats['error']).border = border
                row += 1

        # Adjust column widths
        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 20

    def generate_error_log(
        self,
        results: List[DetectionResult],
        filename: Optional[str] = None
    ) -> str:
        """
        Generate error log file.

        Args:
            results: List of detection results
            filename: Optional custom filename

        Returns:
            Path to generated log file
        """
        if filename is None:
            timestamp = datetime.now().strftime('%Y-%m-%d_%H%M%S')
            filename = f"errors_{timestamp}.log"

        # Put logs in logs directory
        logs_dir = self.output_dir.parent / 'logs'
        logs_dir.mkdir(parents=True, exist_ok=True)
        filepath = logs_dir / filename

        error_results = [r for r in results if r.status == ThinPageStatus.ERROR]

        with open(filepath, 'w') as f:
            f.write(f"# Error Log - Generated {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
            f.write(f"# Total Errors: {len(error_results)}\n\n")

            for result in error_results:
                timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                f.write(f"{timestamp} - ERROR - Failed to analyze {result.url} - {result.error}\n")

        logger.info(f"Generated error log: {filepath}")
        return str(filepath)

    def generate_all_reports(
        self,
        results: List[DetectionResult]
    ) -> Dict[str, str]:
        """
        Generate all report formats.

        Args:
            results: List of detection results

        Returns:
            Dictionary mapping report type to file path
        """
        timestamp = datetime.now().strftime('%Y-%m-%d_%H%M%S')

        reports = {
            'csv': self.generate_csv(
                results,
                f"thin_pages_report_{timestamp}.csv"
            ),
            'excel': self.generate_excel(
                results,
                f"thin_pages_report_{timestamp}.xlsx"
            ),
            'error_log': self.generate_error_log(
                results,
                f"errors_{timestamp}.log"
            )
        }

        return reports


def print_summary(results: List[DetectionResult]):
    """Print summary statistics to console."""
    summary = summarize_results(results)

    print("\n" + "=" * 50)
    print("ANALYSIS COMPLETE")
    print("=" * 50)
    print(f"Total URLs Analyzed: {summary['total']:,}")
    print(f"Thin Pages Found: {summary['thin']:,} ({summary['thin_percent']}%)")
    print(f"Acceptable Pages: {summary['ok']:,} ({summary['ok_percent']}%)")
    print(f"Failed URLs: {summary['error']:,} ({summary['error_percent']}%)")

    if summary.get('language_breakdown'):
        print("\nBy Language:")
        for lang, stats in sorted(summary['language_breakdown'].items()):
            thin_pct = (stats['thin'] / stats['total'] * 100) if stats['total'] > 0 else 0
            print(f"  {lang}: {stats['total']:,} total, {stats['thin']:,} thin ({thin_pct:.1f}%)")

    if summary.get('reason_breakdown'):
        print("\nTop Reasons for Thin Classification:")
        for reason, count in sorted(summary['reason_breakdown'].items(), key=lambda x: -x[1])[:5]:
            print(f"  - {reason}: {count:,}")

    print("=" * 50 + "\n")
